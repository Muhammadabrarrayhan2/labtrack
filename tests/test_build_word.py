import importlib.util
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("build_word", ROOT / "scripts" / "docs" / "build_word.py")
build_word = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build_word)

SOURCE = """title: Dokumen Uji
file: dokumen-uji
numbered: yes
meta: **Penulis:** contoh
---
# !Ringkasan

Paragraf dengan `kode` dan **tebal**.

# Pendahuluan

## Latar belakang

- butir satu
1. langkah satu

Tabel: Data contoh
| Kolom A | Kolom B |
|---|---|
| a1 | b1 |
| a2 | b2 |

> Catatan penting.

```
echo halo
```
"""


def test_front_matter_is_parsed():
    meta, body = build_word.parse_front_matter(SOURCE.splitlines())
    assert meta["title"] == "Dokumen Uji"
    assert meta["meta"] == ["**Penulis:** contoh"]
    assert body[0] == "# !Ringkasan"


def test_build_numbers_headings_and_tables(tmp_path, monkeypatch):
    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "uji.md").write_text(SOURCE, encoding="utf-8")
    monkeypatch.setattr(build_word, "SRC", tmp_path / "src")
    monkeypatch.setattr(build_word, "OUT", tmp_path)
    monkeypatch.setattr(build_word, "ROOT", tmp_path)

    build_word.build("uji")

    doc = Document(tmp_path / "dokumen-uji.docx")
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Ringkasan", "1.  Pendahuluan", "1.1  Latar belakang"]
    assert len(doc.tables) == 1
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Kolom A", "Kolom B"]
    assert any(p.text.startswith("Tabel 1.") for p in doc.paragraphs)
    assert any("echo halo" in p.text for p in doc.paragraphs)
